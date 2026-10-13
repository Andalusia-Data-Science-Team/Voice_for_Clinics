"""
tests/test_eval_integration.py
==============================
End-to-end integration test for Phase 4:
- Synthesizes an Excel workbook mirroring the clinical structure (Scripts, ER Scenarios, OPD Scenarios)
- Imports cases via /eval/gold-cases/import-xlsx endpoint
- Lists imported gold cases via /eval/gold-cases
- Evaluates a pipeline session via /eval/run/{case_id}
- Aggregates report via /eval/report
- Retrieves evaluation runs via /eval/runs
"""
from __future__ import annotations
import openpyxl
import pytest
from fastapi.testclient import TestClient
from mongomock_motor import AsyncMongoMockClient

from app import app
from db.connection import _ensure_indexes
from db.repositories import SessionRepository, GoldDatasetRepository, EvalRepository
from schemas.output_schemas import PipelineSessionOutput


@pytest.fixture
def sample_excel_file(tmp_path):
    """Generate a valid mock Excel workbook matching the schema of Expected_Fields_Expanded.xlsx."""
    wb = openpyxl.Workbook()

    # 1. Scripts sheet
    ws_scripts = wb.active
    ws_scripts.title = "Scripts"
    ws_scripts.append([
        "Scenario ID",
        "Department",
        "Mode",
        "Language",
        "Specialty",
        "Patient Code",
        "Dictation Script (identical for both History OFF and History ON runs)",
    ])
    ws_scripts.append([
        "ER-1",
        "er",
        "doctor",
        "ar",
        "Cardiology",
        "10293",
        "مريض يبلغ من العمر 55 سنة يعاني من ألم حاد ومفاجئ في الصدر",
    ])
    ws_scripts.append([
        "OPD-1",
        "opd",
        "conversation",
        "en",
        "Pulmonology",
        "88412",
        "Patient states they have been wheezing for the past four days.",
    ])

    # 2. ER Scenarios sheet
    ws_er = wb.create_sheet(title="ER Scenarios")
    ws_er.append(["ER-1 — Acute Chest Pain"])
    ws_er.append(["Field", "Source", "Expected Value (History OFF)", "Expected Value (History ON)", "", "Notes"])
    ws_er.append(["Chief Complaint", "Transcript", "ألم حاد في الصدر", "ألم حاد في الصدر", "", ""])
    ws_er.append(["Triage Priority", "Transcript", "Level 1", "Level 1", "", ""])
    ws_er.append(["ICD-10 Codes", "Transcript", "I20.0, I21.9", "I20.0, I21.9", "", ""])
    ws_er.append(["Final Diagnosis", "Transcript", "Acute Myocardial Infarction", "Acute Myocardial Infarction", "", ""])

    # 3. OPD Scenarios sheet
    ws_opd = wb.create_sheet(title="OPD Scenarios")
    ws_opd.append(["OPD-1 — Asthma Exacerbation"])
    ws_opd.append(["Field", "Source", "Expected Value (History OFF)", "Expected Value (History ON)", "", "Notes"])
    ws_opd.append(["Chief Complaint", "Transcript", "Wheezing for four days", "Wheezing for four days", "", ""])
    ws_opd.append(["Final Diagnosis", "Transcript", "Asthma exacerbation", "Asthma exacerbation", "", ""])

    file_path = tmp_path / "test_gold_dataset.xlsx"
    wb.save(file_path)
    return file_path


@pytest.mark.asyncio
async def test_full_evaluation_api_flow(sample_excel_file):
    # Setup test in-memory mock mongo database
    client = AsyncMongoMockClient()
    db = client["medvoice_eval_flow_test"]
    await _ensure_indexes(db)

    app.state.db = db
    app.state.sessions = SessionRepository(db)
    app.state.feedback_repo = None

    test_client = TestClient(app)

    # 1. Import from Excel
    res_import = test_client.post(f"/eval/gold-cases/import-xlsx?file_path={sample_excel_file}")
    assert res_import.status_code == 200
    data_import = res_import.json()
    assert data_import["imported"] == 2

    # 2. List gold cases
    res_list = test_client.get("/eval/gold-cases")
    assert res_list.status_code == 200
    data_list = res_list.json()
    assert data_list["count"] == 2
    case_ids = {c["case_id"] for c in data_list["cases"]}
    assert "ER-1" in case_ids
    assert "OPD-1" in case_ids

    # 3. Create mock completed pipeline session in DB for ER-1
    await app.state.sessions.save_session(
        PipelineSessionOutput(
            session_id="sess-run-er-1",
            department="er",
            language="ar",
            mode="doctor",
            transcript="مريض يبلغ من العمر 55 سنة يعاني من ألم حاد ومفاجئ في الصدر",
            extracted_features={
                "chief_complaint": "ألم حاد في الصدر",
                "triage_priority": "Level 1",
            },
            diagnosis={
                "final_diagnosis": "Acute Myocardial Infarction",
                "icd10_codes": ["I20.0"],
            },
        )
    )

    # 4. Run evaluation endpoint
    res_eval = test_client.post("/eval/run/ER-1?session_id=sess-run-er-1")
    assert res_eval.status_code == 200
    eval_result = res_eval.json()
    assert eval_result["case_id"] == "ER-1"
    assert eval_result["session_id"] == "sess-run-er-1"
    assert eval_result["extraction_accuracy"] >= 0.80
    assert eval_result["diagnosis_match"] is True
    assert eval_result["total_score"] >= 0.85

    # 5. Get aggregate report
    res_report = test_client.get("/eval/report?department=er")
    assert res_report.status_code == 200
    report = res_report.json()
    assert report["total_runs"] == 1
    assert report["diagnosis_match_rate"] == 1.0
    assert report["mean_accuracy"] >= 0.80

    # 6. Retrieve runs
    res_runs = test_client.get("/eval/runs")
    assert res_runs.status_code == 200
    runs = res_runs.json()
    assert runs["count"] == 1
    assert runs["runs"][0]["case_id"] == "ER-1"
