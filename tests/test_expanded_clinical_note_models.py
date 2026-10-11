"""Evaluate all configured models against the expanded clinical-note workbook.

This is a separate opt-in live evaluation for
``Medical_Voice_Assistant_Expected_Fields_Expanded.xlsx``. Every scenario is
run twice (History OFF and History ON), except ``OPD-4``, which is always
ignored as requested. Extraction and reasoning models are evaluated separately;
patient history is supplied to the reasoning/diagnosis call, which is the
pipeline stage that consumes history.

Each output row contains the expected value, model output, score, justification,
and judge model. Values such as ``Not specified`` and an empty model output are
blank-equivalent and score 100 when both mean that no information was provided.

Usage::

    RUN_EXPANDED_EVAL=1 PYTHONPATH=src \
        .venv/bin/python -m pytest tests/test_expanded_clinical_note_models.py -s

Run one scenario first::

    RUN_EXPANDED_EVAL=1 EXPANDED_EVAL_SCENARIO=ER-1 \
        PYTHONPATH=src .venv/bin/python -m pytest \
        tests/test_expanded_clinical_note_models.py -s

Useful variables:

* ``EXPANDED_EVAL_XLSX``: workbook path override.
* ``EXPANDED_EVAL_REPORT``: output XLSX path.
* ``EXPANDED_EVAL_SCENARIO``: one or comma-separated scenario IDs.
* ``EXPANDED_EVAL_HISTORY_ON``: set to ``0`` to run only History OFF.
* ``EXPANDED_EVAL_HISTORY_CSV``: local visit-history fixture override.
* ``EXPANDED_EVAL_FAIL_ON_ERROR``: fail pytest when a model call fails.
"""

from __future__ import annotations

import asyncio
import csv
import json
import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.model_registry import (
    EXTRACTION_MODELS,
    REASONING_MODELS,
    ModelSpec,
    judge_for,
)
from llm.runtime import LLMRuntime
from model.llm_service import LLMService
from model.department_fields import get_department_fields

# Reuse the tested judge parser and blank-equivalence rules without duplicating
# the response-format handling used by the real-use-case evaluator.
from tests import test_real_use_case_models as base_eval

WORKBOOK_PATH = Path(os.environ.get(
    "EXPANDED_EVAL_XLSX",
    "src/data/Medical_Voice_Assistant_Expected_Fields_Expanded.xlsx",
))
REPORT_PATH = Path(os.environ.get(
    "EXPANDED_EVAL_REPORT",
    "reports/expanded_clinical_note_model_evaluation.xlsx",
))
RUN_EVALUATION = os.environ.get("RUN_EXPANDED_EVAL", "").lower() in {
    "1", "true", "yes",
}
HISTORY_ON = os.environ.get("EXPANDED_EVAL_HISTORY_ON", "1").lower() not in {
    "0", "false", "no",
}
FAIL_ON_ERROR = os.environ.get("EXPANDED_EVAL_FAIL_ON_ERROR", "").lower() in {
    "1", "true", "yes",
}
SCENARIO_FILTER = {
    item.strip().upper()
    for item in os.environ.get("EXPANDED_EVAL_SCENARIO", "").split(",")
    if item.strip()
}
HISTORY_CSV = Path(os.environ.get(
    "EXPANDED_EVAL_HISTORY_CSV",
    "src/data/Visit Level History.csv",
))

SCENARIO_SHEETS = ("ER Scenarios", "OPD Scenarios", "Dental Scenarios")
SKIPPED_SCENARIOS = {"OPD-4"}

FIELD_ALIASES = {
    "plan_of_care_short_long_term": "plan_of_care",
    "triage_classification": "triage_priority",
}
FIELD_DISPLAY_NAMES = {
    "plan_of_care": "Plan of Care",
    "icd10_codes": "ICD-10 Codes",
    "final_diagnosis": "Final Diagnosis",
    "final_diagnosis_rationale": "Final Diagnosis Rationale",
    "triage_priority": "Triage Priority",
    "differential_diagnoses": "Differential Diagnoses",
    "history_considered": "History Considered",
    "history_notes": "History Notes",
}
REASONING_FIELDS = {
    "icd10_codes",
    "final_diagnosis",
    "final_diagnosis_rationale",
    "triage_priority",
    "differential_diagnoses",
    "history_considered",
    "history_notes",
}


@dataclass(frozen=True)
class ExpandedScenario:
    """One workbook script and its expected OFF/ON field values."""

    scenario_id: str
    department: str
    language: str
    is_conversation: bool
    specialty: str | None
    patient_code: str | None
    transcript: str
    expected_fields: dict[str, tuple[str, str, str]]


@dataclass
class ResultRow:
    """One model/stage/history/field result for the XLSX report."""

    scenario: str
    department: str
    language: str
    phase: str
    stage: str
    model_id: str
    model_label: str
    model_slug: str
    field: str
    expected: str
    output: str
    score: float | None = None
    justification: str = ""
    judge_model: str = ""
    latency_sec: float | None = None
    cost_usd: float | None = None
    prompt_tokens: int | None = None
    completion_tokens: int | None = None
    error: str = ""
    field_in_workbook: bool = True
    source: str = ""


def _text(value: Any) -> str:
    """Serialize model values for an Excel cell."""
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value).strip()


def _norm_label(label: Any) -> str:
    """Normalize a display label for schema matching."""
    return "".join(ch.lower() for ch in str(label or "") if ch.isalnum())


def _key(label: Any, department: str = "") -> str:
    """Normalize workbook labels to extraction JSON keys."""
    normalized = _norm_label(label)
    field_map = {
        _norm_label(item["label"]): item["key"]
        for item in get_department_fields(department)
    }
    field_map.update({
        _norm_label("Triage Priority"): "triage_priority",
        _norm_label("ICD-10 Codes"): "icd10_codes",
        _norm_label("Final Diagnosis"): "final_diagnosis",
        _norm_label("Final Diagnosis Rationale"): "final_diagnosis_rationale",
        _norm_label("Differential Diagnoses"): "differential_diagnoses",
        _norm_label("History Considered"): "history_considered",
        _norm_label("History Notes"): "history_notes",
    })
    if normalized in field_map:
        return field_map[normalized]
    aliases = {
        "chiefcomplaintco": "chief_complaint",
        "icd10codes": "icd10_codes",
        "finaldiagnosis": "final_diagnosis",
        "finaldiagnosisrationale": "final_diagnosis_rationale",
        "triageclassification": "triage_priority",
        "triagepriority": "triage_priority",
        "differentialdiagnoses": "differential_diagnoses",
        "historyconsidered": "history_considered",
        "historynotes": "history_notes",
    }
    key = aliases.get(normalized)
    if key is None:
        raw = str(label or "").lower().replace("/", " ").replace("-", " ")
        key = "_".join(raw.replace("(", "").replace(")", "").split())
    return FIELD_ALIASES.get(key, key)


def _display(key: str) -> str:
    """Return one unified report label."""
    return FIELD_DISPLAY_NAMES.get(key, key.replace("_", " ").title())


def _scenario_scripts(workbook: openpyxl.Workbook) -> dict[str, dict[str, Any]]:
    """Read scripts and metadata from the expanded Scripts sheet."""
    rows = list(workbook["Scripts"].iter_rows(values_only=True))
    header = {str(value).strip(): i for i, value in enumerate(rows[0])}
    output: dict[str, dict[str, Any]] = {}
    for row in rows[1:]:
        if not row or not row[0]:
            continue
        scenario_id = str(row[header["Scenario"]]).strip().upper()
        if scenario_id in SKIPPED_SCENARIOS:
            continue
        mode = str(row[header["Mode"]] or "").lower()
        language = str(row[header["Language"]] or "").lower()
        specialty = str(row[header["Specialty"]] or "").strip()
        patient_code = str(row[header["Patient Code"]] or "").split()[0] or None
        output[scenario_id] = {
            "department": str(row[header["Department"]]).lower(),
            "language": "ar" if language.startswith("ar") else "en",
            "is_conversation": mode == "conversation",
            "specialty": None if specialty.lower().startswith("general") else specialty,
            "patient_code": patient_code,
            "transcript": str(row[header["Dictation Script (identical for both History OFF and History ON runs)"]] or ""),
        }
    return output


def _scenario_fields(
    workbook: openpyxl.Workbook,
    departments: dict[str, str],
) -> dict[str, dict[str, tuple[str, str, str]]]:
    """Read Field/Source/Expected-OFF/Expected-ON tables."""
    result: dict[str, dict[str, tuple[str, str, str]]] = {}
    for sheet_name in SCENARIO_SHEETS:
        current = ""
        for row in workbook[sheet_name].iter_rows(values_only=True):
            first = str(row[0] or "").strip() if row else ""
            match = re.match(r"^([A-Za-z]+-\d+)\b", first)
            if first and len(row) > 1 and row[1] is None and match:
                current = match.group(1).upper()
                result.setdefault(current, {})
                continue
            if first == "Field" or not current or len(row) < 4 or not row[1]:
                continue
            key = _key(first, departments.get(current, "opd"))
            source = str(row[1]).strip()
            off = "" if row[2] is None else str(row[2]).strip()
            on = "" if row[3] is None else str(row[3]).strip()
            if key in result[current] and result[current][key][1:] != (off, on):
                old_source, old_off, old_on = result[current][key]
                off = " / ".join(value for value in (old_off, off) if value)
                on = " / ".join(value for value in (old_on, on) if value)
                source = " / ".join(value for value in (old_source, source) if value)
            result[current][key] = (source, off, on)
    return result


def load_scenarios(path: Path = WORKBOOK_PATH) -> list[ExpandedScenario]:
    """Load all expanded scenarios except OPD-4."""
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    scripts = _scenario_scripts(workbook)
    departments = {
        scenario_id: meta["department"]
        for scenario_id, meta in scripts.items()
    }
    fields = _scenario_fields(workbook, departments)
    scenarios = []
    for scenario_id, meta in scripts.items():
        if SCENARIO_FILTER and scenario_id not in SCENARIO_FILTER:
            continue
        scenarios.append(ExpandedScenario(
            scenario_id=scenario_id,
            department=meta["department"],
            language=meta["language"],
            is_conversation=meta["is_conversation"],
            specialty=meta["specialty"],
            patient_code=meta["patient_code"],
            transcript=meta["transcript"],
            expected_fields=fields.get(scenario_id, {}),
        ))
    return sorted(scenarios, key=lambda item: item.scenario_id)


def _history_visits(patient_code: str | None) -> list[dict[str, str]]:
    """Load workbook fixture history using its tab-separated local format."""
    if not patient_code or not HISTORY_CSV.exists():
        return []
    for encoding in ("cp1252", "latin1", "utf-8"):
        try:
            with HISTORY_CSV.open(encoding=encoding, newline="") as handle:
                rows = csv.DictReader(handle, delimiter="\t")
                return [
                    {
                        "visit_id": row.get("VisitID", ""),
                        "visit_date": row.get("VisitDate", ""),
                        "specialty": row.get("Specialty", ""),
                        "icd10_diagnoses": [
                            item.strip() for item in (row.get("CurrentICD10") or "").split("|")
                            if item.strip() and item.strip().upper() != "NULL"
                        ],
                        "department": row.get("Department", ""),
                        "visit_type": row.get("VisitTypeName", ""),
                        "previous_visit_count": row.get("PreviousVisitCount", "0"),
                    }
                    for row in rows
                    if row.get("PatientCode") == patient_code
                ]
        except UnicodeDecodeError:
            continue
    return []


def _history_prompt(visits: list[dict[str, str]]) -> str:
    """Format history in the same compact style used by the application."""
    if not visits:
        return ""
    return "\n".join(
        f"- {v['visit_date']} | {v['specialty']} | {v['department']} "
        f"({v['visit_type']}): {('; '.join(v['icd10_diagnoses']) or 'No diagnosis recorded')}"
        for v in visits
    )


def _field_bucket(source: str, expected: str) -> str:
    """Classify an expanded-workbook field before calling the judge."""
    source_text = (source or "").lower()
    empty_markers = (
        "not mentioned",
        "not applicable",
        "n/a",
        "not explicitly asked",
        "not clearly stated",
        "not reported",
        "not specified",
    )
    if any(marker in source_text for marker in empty_markers):
        return "empty"
    if "ambiguous" in source_text:
        return "ambiguous"
    if base_eval._is_blank_like(expected):
        return "empty"
    return "standard"


ICD10_CODE_RE = re.compile(r"\b[A-Z][0-9]{2}(?:\.[0-9A-Z]{1,4})?\b")


def _icd10_codes(value: str) -> set[str]:
    """Extract ICD-10-like codes from text or serialized model output."""
    return set(ICD10_CODE_RE.findall((value or "").upper()))


def _same_icd10_parent(expected: str, actual: str) -> bool:
    """Return whether expected and actual share a letter-plus-two-digit parent."""
    expected_parents = {code[:3] for code in _icd10_codes(expected)}
    actual_parents = {code[:3] for code in _icd10_codes(actual)}
    return bool(expected_parents & actual_parents)


def _is_icd10_field(field: str) -> bool:
    """Identify report fields whose score includes the ICD-10 baseline."""
    return field in {"ICD-10 Codes", "Final Diagnosis"}


def _make_rows(
    scenario: ExpandedScenario,
    phase: str,
    stage: str,
    model: ModelSpec,
    output: dict[str, Any] | None,
    meta: dict[str, Any] | None,
    error: str = "",
) -> list[ResultRow]:
    """Create rows only for fields present in the workbook."""
    data = (output or {}).get("json_data", {}) if isinstance(output, dict) else {}
    selected = {
        key: expected for key, expected in scenario.expected_fields.items()
        if (stage == "Reasoning") == (key in REASONING_FIELDS)
    }
    return [
        ResultRow(
            scenario=scenario.scenario_id,
            department=scenario.department,
            language=scenario.language,
            phase=phase,
            stage=stage,
            model_id=model.id,
            model_label=model.label,
            model_slug=model.slug,
            field=_display(key),
            expected=expected[1] if phase == "OFF" else expected[2],
            output=_text(data.get(key)),
            latency_sec=(meta or {}).get("latency_sec"),
            cost_usd=(meta or {}).get("cost_usd"),
            prompt_tokens=(meta or {}).get("prompt_tokens"),
            completion_tokens=(meta or {}).get("completion_tokens"),
            source=expected[0],
            error=error,
        )
        for key, expected in selected.items()
    ]


async def _score_rows(runtime: LLMRuntime, rows: list[ResultRow]) -> None:
    """Assign deterministic blank scores or cross-family judge scores."""
    candidates = {
        (model.id, model.slug): model
        for model in (*EXTRACTION_MODELS.values(), *REASONING_MODELS.values())
    }
    groups: dict[tuple[str, str, str, str], list[ResultRow]] = {}
    for row in rows:
        groups.setdefault((row.scenario, row.phase, row.stage, row.model_slug), []).append(row)

    for (scenario_id, phase, stage, model_slug), group in groups.items():
        candidate = next((m for (mid, slug), m in candidates.items() if slug == model_slug), None)
        if candidate is None:
            continue
        judge = judge_for(candidate)
        judge_rows: list[base_eval.EvaluationRow] = []
        judgeable_rows: list[ResultRow] = []
        for row in group:
            row.judge_model = judge.label
            if row.error:
                row.justification = "Not scored because the model call failed."
                continue
            bucket = _field_bucket(row.source, row.expected)
            if bucket == "empty" or base_eval._is_blank_like(row.expected):
                row.score = 100.0 if base_eval._is_blank_like(row.output) else 20.0
                row.justification = (
                    "Blank-equivalent value as expected."
                    if row.score == 100.0
                    else "Expected blank-equivalent value, but model populated it."
                )
                continue
            if bucket == "ambiguous" and base_eval._is_blank_like(row.output):
                row.score = 100.0
                row.justification = "Blank is acceptable for this ambiguous field."
                continue
            if base_eval._is_blank_like(row.output):
                row.score = 0.0
                row.justification = "Missing expected content."
                continue
            judge_rows.append(base_eval.EvaluationRow(
                scenario_id=row.scenario,
                department=row.department,
                language=row.language,
                stage=row.stage,
                model_id=candidate.id,
                model_label=candidate.label,
                model_slug=candidate.slug,
                field=row.field,
                expected=row.expected,
                output=row.output,
                latency_sec=None,
                cost_usd=None,
                prompt_tokens=None,
                completion_tokens=None,
            ))
            judgeable_rows.append(row)
        if judge_rows:
            scored = await base_eval._judge_output(
                runtime, judge, scenario_id, stage, judge_rows
            )
            # Match by the field name, never by list position. Deterministic
            # blank/ambiguous rows are omitted from judge_rows, so zip() could
            # attach a judge result to the wrong report row.
            for row in judgeable_rows:
                judged_score, justification = scored.get(
                    row.field, (None, "Judge omitted this field.")
                )
                if judged_score is None:
                    row.score = None
                    row.justification = justification
                    continue
                if _is_icd10_field(row.field) and _same_icd10_parent(
                    row.expected, row.output
                ):
                    row.score = 70.0 + (0.30 * judged_score)
                    row.justification = (
                        "Shared ICD-10 parent code; 70% baseline plus semantic "
                        f"judge score ({judged_score:.0f}%). {justification}"
                    )
                else:
                    row.score = judged_score
                    row.justification = justification


def _write_report(rows: list[ResultRow], path: Path) -> None:
    """Write detailed results, summary, and errors to an XLSX report."""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    details = workbook.active
    details.title = "Detailed Results"
    headers = [
        "Scenario", "Department", "Language", "Phase", "Stage", "Model ID",
        "Model", "Model Slug", "Field", "Expected", "Output", "Score (0-100)",
        "Justification", "Judge Model", "Latency (sec)", "Cost (USD)",
        "Prompt Tokens", "Completion Tokens", "Error",
    ]
    details.append(headers)
    for row in rows:
        details.append([
            row.scenario, row.department, row.language, row.phase, row.stage,
            row.model_id, row.model_label, row.model_slug, row.field, row.expected,
            row.output, row.score, row.justification, row.judge_model,
            row.latency_sec, row.cost_usd, row.prompt_tokens,
            row.completion_tokens, row.error,
        ])

    summary = workbook.create_sheet("Summary")
    summary.append(["Scenario", "Phase", "Stage", "Model", "Average Score", "Latency (sec)", "Cost (USD)", "Errors"])
    grouped: dict[tuple[str, str, str, str], list[ResultRow]] = {}
    for row in rows:
        grouped.setdefault((row.scenario, row.phase, row.stage, row.model_label), []).append(row)
    for key, values in grouped.items():
        scores = [row.score for row in values if row.score is not None]
        latencies = [row.latency_sec for row in values if row.latency_sec is not None]
        costs = [row.cost_usd for row in values if row.cost_usd is not None]
        errors = " | ".join(sorted({row.error for row in values if row.error}))
        summary.append([*key, sum(scores) / len(scores) if scores else None, max(latencies, default=None), max(costs, default=None), errors])

    errors_sheet = workbook.create_sheet("Errors")
    errors_sheet.append(headers)
    for row in rows:
        if row.error:
            errors_sheet.append([row.scenario, row.department, row.language, row.phase, row.stage, row.model_id, row.model_label, row.model_slug, row.field, row.expected, row.output, row.score, row.justification, row.judge_model, row.latency_sec, row.cost_usd, row.prompt_tokens, row.completion_tokens, row.error])

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="6B4F35")
            cell.alignment = Alignment(horizontal="center")
        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            sheet.column_dimensions[letter].width = min(55, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)


@pytest.mark.llm
@pytest.mark.asyncio
@pytest.mark.skipif(
    not RUN_EVALUATION,
    reason="Set RUN_EXPANDED_EVAL=1 to run the live expanded evaluation",
)
async def test_expanded_clinical_note_models_generate_report() -> None:
    """Evaluate extraction and reasoning models for History OFF and ON."""
    if not WORKBOOK_PATH.exists():
        pytest.fail(f"Expanded workbook not found: {WORKBOOK_PATH}")
    scenarios = load_scenarios()
    if not scenarios:
        pytest.fail("No expanded scenarios matched the requested filter")

    rows: list[ResultRow] = []
    runtime = LLMRuntime()
    try:
        for scenario in scenarios:
            refined = await _refine_and_translate(scenario, runtime)
            phases = [("OFF", "")]
            if HISTORY_ON:
                phases.append(("ON", _history_prompt(_history_visits(scenario.patient_code))))
            for phase, history_text in phases:
                for model in EXTRACTION_MODELS.values():
                    meta: dict[str, Any] = {}
                    try:
                        output = await LLMService.extract_features_async(
                            translated_text=refined,
                            runtime=runtime,
                            department=scenario.department,
                            is_conversation=scenario.is_conversation,
                            model_id=model.id,
                            meta=meta,
                        )
                        rows.extend(_make_rows(scenario, phase, "Extraction", model, output, meta))
                    except Exception as exc:
                        rows.extend(_make_rows(scenario, phase, "Extraction", model, None, meta, str(exc)))
                for model in REASONING_MODELS.values():
                    meta = {}
                    try:
                        output = await LLMService.recommend_diagnosis_async(
                            translated_text=refined,
                            runtime=runtime,
                            department=scenario.department,
                            patient_history_text=history_text,
                            is_conversation=scenario.is_conversation,
                            model_id=model.id,
                            meta=meta,
                        )
                        rows.extend(_make_rows(scenario, phase, "Reasoning", model, output, meta))
                    except Exception as exc:
                        rows.extend(_make_rows(scenario, phase, "Reasoning", model, None, meta, str(exc)))
    finally:
        await runtime.close()

    judge_runtime = LLMRuntime()
    try:
        await _score_rows(judge_runtime, rows)
    finally:
        await judge_runtime.close()
    _write_report(rows, REPORT_PATH)
    print(f"Wrote {len(rows)} rows to {REPORT_PATH}")
    errors = sorted({row.error for row in rows if row.error})
    if FAIL_ON_ERROR and errors:
        pytest.fail("Expanded evaluation errors:\n" + "\n".join(errors))


async def _refine_and_translate(
    scenario: ExpandedScenario,
    runtime: LLMRuntime,
) -> str:
    """Run the shared refinement/translation pipeline once per scenario."""
    if scenario.language == "ar":
        stream = LLMService.refine_ar_transcription_stream_async(
            scenario.transcript,
            runtime,
            is_conversation=scenario.is_conversation,
            specialty=scenario.specialty,
        )
    else:
        stream = LLMService.refine_en_transcription_stream_async(
            scenario.transcript,
            runtime,
            is_conversation=scenario.is_conversation,
            specialty=scenario.specialty,
        )
    refined = "".join([chunk async for chunk in stream])
    if scenario.language == "ar":
        return await LLMService.translate_to_eng_async(
            refined, runtime, is_conversation=scenario.is_conversation
        )
    return refined
