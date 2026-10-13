"""
evaluation/importers.py
=======================
Importer for Gold Cases from Excel workbooks (e.g. Medical_Voice_Assistant_Expected_Fields_Expanded.xlsx).
Maps multi-tab scenarios into validated GoldCase objects.
"""
from __future__ import annotations
import os
import re
import logging
from pathlib import Path
from typing import Any, Dict, List, Optional
import openpyxl

from schemas.gold_schemas import GoldCase
from model.department_fields import get_department_fields

logger = logging.getLogger(__name__)

SCENARIO_SHEETS = ["ER Scenarios", "OPD Scenarios", "Dental Scenarios"]
SCENARIO_ID_RE = re.compile(r"^([A-Za-z]+-\d+)\b")
PATIENT_CODE_RE = re.compile(r"\d+")

EXTRA_FIELDS = {
    "triage_priority": "Triage Priority",
    "icd10_codes": "ICD-10 Codes",
    "final_diagnosis": "Final Diagnosis",
    "final_diagnosis_rationale": "Final Diagnosis Rationale",
    "history_considered": "history_considered",
    "history_notes": "history_notes",
}

EXCLUDED_FIELDS = {"final_diagnosis_rationale", "history_considered", "history_notes"}


def _norm_label(label: str) -> str:
    """Normalize labels for robust matching."""
    return re.sub(r"[^a-z0-9]", "", (label or "").lower())


def _build_label_key_map(department: str) -> Dict[str, str]:
    """Map human labels in the Excel sheet to extraction dictionary keys."""
    m: Dict[str, str] = {}
    try:
        fields = get_department_fields(department)
        for f in fields:
            m[_norm_label(f["label"])] = f["key"]
    except Exception as exc:
        logger.warning("[Importer] Failed to load department fields for %s: %s", department, exc)

    for key, label in EXTRA_FIELDS.items():
        m[_norm_label(label)] = key
    return m


def _parse_scripts(wb: openpyxl.Workbook) -> Dict[str, Dict[str, Any]]:
    """Parse 'Scripts' sheet containing scenario metadata and scripts."""
    if "Scripts" not in wb.sheetnames:
        return {}

    ws = wb["Scripts"]
    rows = list(ws.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        return {}

    header = rows[0]
    idx = {str(name).strip(): i for i, name in enumerate(header) if name is not None}
    out: Dict[str, Dict[str, Any]] = {}

    script_col = "Dictation Script (identical for both History OFF and History ON runs)"
    script_idx = idx.get(script_col)
    if script_idx is None:
        # Fallback to finding script column by substring
        for name, col_idx in idx.items():
            if "script" in name.lower():
                script_idx = col_idx
                break

    for row in rows[1:]:
        if not row or not row[0]:
            continue
        scenario_id = str(row[0]).strip()
        dept = str(row[idx["Department"]]).strip().lower() if "Department" in idx else "opd"
        mode = str(row[idx["Mode"]]).strip().lower() if "Mode" in idx else "doctor"
        lang_raw = str(row[idx["Language"]]).strip().lower() if "Language" in idx else "ar"
        specialty = str(row[idx["Specialty"]]).strip() if "Specialty" in idx and row[idx["Specialty"]] else None
        if specialty and specialty.lower().startswith("general"):
            specialty = None

        script = str(row[script_idx]).strip() if script_idx is not None and row[script_idx] else ""

        out[scenario_id] = {
            "department": dept,
            "mode": mode,
            "language": "ar" if lang_raw.startswith("ar") else "en",
            "specialty": specialty,
            "script": script,
        }
    return out


def _parse_scenario_sheet(ws: openpyxl.worksheet.worksheet.Worksheet) -> Dict[str, Dict[str, str]]:
    """Parse scenario sub-tables from ER/OPD/Dental sheets."""
    rows = list(ws.iter_rows(values_only=True))
    scenarios: Dict[str, Dict[str, str]] = {}
    current_id: Optional[str] = None

    for row in rows:
        if not row:
            continue
        first = row[0]
        # Title row: e.g. "ER-1 — Nose Wound Recurrence (...)"
        if isinstance(first, str) and (len(row) <= 1 or row[1] is None):
            m = SCENARIO_ID_RE.match(first.strip())
            if m:
                current_id = m.group(1)
                scenarios.setdefault(current_id, {})
            continue

        if first == "Field" and len(row) > 1 and row[1] == "Source":
            continue

        if current_id and first and len(row) > 2 and row[1]:
            label = str(first).strip()
            # row[2] = expected_off, row[3] = expected_on (prefer non-empty value)
            exp_off = "" if row[2] is None else str(row[2]).strip()
            exp_on = "" if len(row) <= 3 or row[3] is None else str(row[3]).strip()
            chosen_val = exp_off if exp_off else exp_on
            scenarios[current_id][label] = chosen_val

    return scenarios


def load_gold_cases_from_xlsx(file_path: str | Path) -> List[GoldCase]:
    """Parse an Excel workbook into a list of validated GoldCase models."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Gold dataset Excel file not found: {path}")

    wb = openpyxl.load_workbook(path, data_only=True)
    scripts = _parse_scripts(wb)

    sheet_scenarios: Dict[str, Dict[str, str]] = {}
    for name in SCENARIO_SHEETS:
        if name in wb.sheetnames:
            sheet_scenarios.update(_parse_scenario_sheet(wb[name]))

    gold_cases: List[GoldCase] = []

    for scenario_id, meta in scripts.items():
        script_text = meta.get("script", "")
        if not script_text or len(script_text) < 10:
            continue

        department = meta.get("department", "opd").lower()
        if department not in ("er", "opd", "dental"):
            department = "opd"

        label_map = _build_label_key_map(department)
        raw_fields = sheet_scenarios.get(scenario_id, {})

        expected_extraction: Dict[str, Any] = {}
        expected_diagnosis: Dict[str, Any] = {}

        for label, val in raw_fields.items():
            norm_key = label_map.get(_norm_label(label))
            if not norm_key or norm_key in EXCLUDED_FIELDS:
                continue

            if norm_key == "icd10_codes":
                codes = [c.strip() for c in str(val).split(",") if c.strip()]
                expected_diagnosis["icd10_codes"] = codes
            elif norm_key == "final_diagnosis":
                expected_diagnosis["final_diagnosis"] = val
            else:
                expected_extraction[norm_key] = val

        gold_case = GoldCase(
            case_id=scenario_id,
            department=department,
            specialty=meta.get("specialty"),
            audio_language=meta.get("language", "ar"),
            transcript=script_text,
            expected_extraction=expected_extraction,
            expected_diagnosis=expected_diagnosis if expected_diagnosis else None,
            source="xlsx_import",
            tags=[department, meta.get("language", "ar")],
        )
        gold_cases.append(gold_case)

    logger.info("[Importer] Successfully parsed %d gold cases from %s", len(gold_cases), path.name)
    return gold_cases
